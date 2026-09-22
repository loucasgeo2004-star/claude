#!/usr/bin/env python3
"""
Excel Formula Re-Pointer for Trident Unity LP
Automatically re-points BS and IS column I formulas from T to TB2026
"""

import openpyxl
import re
import sys
from pathlib import Path
from datetime import datetime


def build_mapping(wb):
    """Build mapping of T cell references to TB2026 cells"""
    ws_tb = wb['TB2026']
    mapping = {}

    for row in range(1, ws_tb.max_row + 1):
        dr = ws_tb[f'D{row}'].value
        cr = ws_tb[f'E{row}'].value

        if dr and isinstance(dr, str) and dr.startswith('='):
            t_ref = dr.replace('=', '').strip()
            if 'T!' in t_ref:
                mapping[t_ref] = f'TB2026!D{row}'

        if cr and isinstance(cr, str) and cr.startswith('='):
            t_ref = cr.replace('=', '').strip()
            if 'T!' in t_ref:
                mapping[t_ref] = f'TB2026!E{row}'

    return mapping


def update_sheet_formulas(ws, mapping, sheet_name):
    """Update formulas in a sheet using the mapping"""
    updates = []

    for row in range(1, ws.max_row + 1):
        cell = ws[f'I{row}']

        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            original = cell.value
            updated = original

            # Replace each T! reference with TB2026 reference
            for t_ref, tb_ref in mapping.items():
                updated = updated.replace(t_ref, tb_ref)

            if updated != original:
                ws[f'I{row}'].value = updated
                updates.append({
                    'cell': f'{sheet_name}!I{row}',
                    'before': original,
                    'after': updated
                })

    return updates


def repoint_formulas(input_file):
    """Main function to re-point formulas in Excel file"""

    # Verify file exists
    if not Path(input_file).exists():
        print(f"❌ Error: File not found: {input_file}")
        sys.exit(1)

    print("="*80)
    print("TRIDENT UNITY LP - FORMULA RE-POINTER")
    print("="*80)
    print(f"\nInput file: {input_file}")
    print(f"Timestamp: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n")

    try:
        # Load workbook
        print("📂 Loading workbook...")
        wb = openpyxl.load_workbook(input_file)

        # Verify required sheets exist
        required_sheets = ['TB2026', 'BS', 'IS', 'T']
        for sheet in required_sheets:
            if sheet not in wb.sheetnames:
                print(f"❌ Error: Required sheet '{sheet}' not found!")
                sys.exit(1)

        print("✓ All required sheets found (TB2026, BS, IS, T)")

        # Build mapping
        print("\n🔍 Building T → TB2026 mapping...")
        mapping = build_mapping(wb)
        print(f"✓ Found {len(mapping)} T cell → TB2026 mappings")

        # Update BS sheet
        print("\n📊 Updating BS (Balance Sheet)...")
        ws_bs = wb['BS']
        bs_updates = update_sheet_formulas(ws_bs, mapping, 'BS')
        print(f"✓ Updated {len(bs_updates)} BS formulas")

        # Update IS sheet
        print("\n📈 Updating IS (Income Statement)...")
        ws_is = wb['IS']
        is_updates = update_sheet_formulas(ws_is, mapping, 'IS')
        print(f"✓ Updated {len(is_updates)} IS formulas")

        # Generate output file name
        input_path = Path(input_file)
        output_file = input_path.parent / f"{input_path.stem}_REPOINTED{input_path.suffix}"

        # Save updated workbook
        print(f"\n💾 Saving updated file...")
        wb.save(output_file)
        print(f"✓ Saved to: {output_file}")

        # Print summary
        print("\n" + "="*80)
        print("SUMMARY")
        print("="*80)
        print(f"Total formulas updated: {len(bs_updates) + len(is_updates)}")
        print(f"  - BS (Balance Sheet): {len(bs_updates)} cells")
        print(f"  - IS (Income Statement): {len(is_updates)} cells")

        if bs_updates or is_updates:
            print("\nSample updates:")
            all_updates = bs_updates + is_updates
            for update in all_updates[:3]:
                print(f"\n  {update['cell']}:")
                print(f"    Before: {update['before'][:60]}")
                print(f"    After:  {update['after'][:60]}")

        print("\n" + "="*80)
        print("✅ DONE! Next steps:")
        print("  1. Open the updated file in Excel/LibreOffice")
        print("  2. Let it recalculate (Ctrl+Shift+F9 or Tools → Recalculate)")
        print("  3. Verify that all values match the original")
        print("="*80 + "\n")

        return True

    except Exception as e:
        print(f"❌ Error: {str(e)}")
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print("Usage: python3 excel_formula_repointer.py <path_to_excel_file>")
        print("\nExample:")
        print("  python3 excel_formula_repointer.py accounting_2026.xlsx")
        sys.exit(1)

    input_file = sys.argv[1]
    repoint_formulas(input_file)
